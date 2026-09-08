"""
Backend API Testing for Semoyo Joyo B2B App - Tahap 2
Tests cost_price, profit/loss, stock management, order edit/delete
"""
import requests
import json
from typing import Optional

BASE_URL = "https://ship-it-demo.preview.emergentagent.com/api"
ADMIN_TOKEN: Optional[str] = None
OWNER_TOKEN: Optional[str] = None

def log(msg: str):
    print(f"[TEST] {msg}")

def error(msg: str):
    print(f"[ERROR] {msg}")

def success(msg: str):
    print(f"[SUCCESS] {msg}")

def admin_login():
    """Login as admin and get access token"""
    global ADMIN_TOKEN
    log("Logging in as admin...")
    resp = requests.post(f"{BASE_URL}/admin/login", json={
        "username": "admin",
        "password": "admin123"
    })
    if resp.status_code != 200:
        error(f"Admin login failed: {resp.status_code} - {resp.text}")
        return False
    data = resp.json()
    ADMIN_TOKEN = data["access_token"]
    success(f"Admin logged in successfully. Token: {ADMIN_TOKEN[:20]}...")
    return True

def owner_login():
    """Login as owner and get access token"""
    global OWNER_TOKEN
    log("Logging in as owner...")
    resp = requests.post(f"{BASE_URL}/admin/login", json={
        "username": "owner",
        "password": "owner123"
    })
    if resp.status_code != 200:
        error(f"Owner login failed: {resp.status_code} - {resp.text}")
        return False
    data = resp.json()
    OWNER_TOKEN = data["access_token"]
    success(f"Owner logged in successfully. Token: {OWNER_TOKEN[:20]}...")
    return True

def get_headers():
    """Get headers with admin token"""
    return {"Authorization": f"Bearer {ADMIN_TOKEN}"}

def get_owner_headers():
    """Get headers with owner token"""
    return {"Authorization": f"Bearer {OWNER_TOKEN}"}

def test_product_cost_price():
    """Test 1: Product cost_price + profit fields"""
    log("\n=== TEST 1: Product cost_price + profit fields ===")
    
    # Get categories first
    resp = requests.get(f"{BASE_URL}/admin/categories", headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to get categories: {resp.status_code}")
        return False
    categories = resp.json()
    if not categories:
        error("No categories found")
        return False
    category_id = categories[0]["id"]
    log(f"Using category: {categories[0]['name']} (ID: {category_id})")
    
    # Test 1a: GET /api/admin/products - should have cost_price, profit_per_unit, margin_pct
    log("Testing GET /api/admin/products (admin view)...")
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to get admin products: {resp.status_code}")
        return False
    admin_products = resp.json()
    if admin_products:
        product = admin_products[0]
        log(f"Sample product: {product['name']}")
        log(f"  - price: {product.get('price')}")
        log(f"  - cost_price: {product.get('cost_price')}")
        log(f"  - profit_per_unit: {product.get('profit_per_unit')}")
        log(f"  - margin_pct: {product.get('margin_pct')}")
        if 'cost_price' not in product or 'profit_per_unit' not in product or 'margin_pct' not in product:
            error("Admin products missing cost_price/profit_per_unit/margin_pct fields")
            return False
    success("Admin products have cost_price, profit_per_unit, margin_pct")
    
    # Test 1b: GET /api/products (public) - cost_price must be 0
    log("Testing GET /api/products (public view)...")
    resp = requests.get(f"{BASE_URL}/products")
    if resp.status_code != 200:
        error(f"Failed to get public products: {resp.status_code}")
        return False
    public_products = resp.json()
    if public_products:
        product = public_products[0]
        log(f"Sample public product: {product['name']}")
        log(f"  - price: {product.get('price')}")
        log(f"  - cost_price: {product.get('cost_price')}")
        if product.get('cost_price', 0) != 0:
            error(f"Public product exposes cost_price: {product.get('cost_price')}")
            return False
    success("Public products do NOT expose cost_price (cost_price=0)")
    
    # Test 1c: POST /api/admin/products with cost_price
    log("Testing POST /api/admin/products with cost_price...")
    new_product = {
        "category_id": category_id,
        "name": "Test Beras Premium Organik",
        "description": "Beras organik berkualitas tinggi untuk testing",
        "price": 10000,
        "cost_price": 7000,
        "unit": "kg",
        "min_order": 5,
        "stock": 50,
        "is_active": True
    }
    resp = requests.post(f"{BASE_URL}/admin/products", json=new_product, headers=get_headers())
    if resp.status_code != 201:
        error(f"Failed to create product: {resp.status_code} - {resp.text}")
        return False
    created = resp.json()
    product_id = created["id"]
    log(f"Created product: {created['name']} (ID: {product_id})")
    log(f"  - price: {created['price']}")
    log(f"  - cost_price: {created['cost_price']}")
    log(f"  - profit_per_unit: {created['profit_per_unit']}")
    log(f"  - margin_pct: {created['margin_pct']}")
    
    # Verify calculations
    expected_profit = 10000 - 7000
    expected_margin = 30.0
    if abs(created['profit_per_unit'] - expected_profit) > 0.01:
        error(f"Profit calculation wrong: expected {expected_profit}, got {created['profit_per_unit']}")
        return False
    if abs(created['margin_pct'] - expected_margin) > 0.1:
        error(f"Margin calculation wrong: expected {expected_margin}%, got {created['margin_pct']}%")
        return False
    success(f"Product created with correct profit_per_unit={expected_profit} and margin_pct={expected_margin}%")
    
    # Test 1d: PUT /api/admin/products - update cost_price and stock
    log("Testing PUT /api/admin/products - update cost_price and stock...")
    updated_product = {
        "category_id": category_id,
        "name": "Test Beras Premium Organik",
        "description": "Beras organik berkualitas tinggi untuk testing",
        "price": 10000,
        "cost_price": 8000,  # Changed from 7000 to 8000
        "unit": "kg",
        "min_order": 5,
        "stock": 60,  # Changed from 50 to 60
        "is_active": True
    }
    resp = requests.put(f"{BASE_URL}/admin/products/{product_id}", json=updated_product, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to update product: {resp.status_code} - {resp.text}")
        return False
    updated = resp.json()
    log(f"Updated product: {updated['name']}")
    log(f"  - cost_price: {updated['cost_price']} (was 7000)")
    log(f"  - stock: {updated['stock']} (was 50)")
    log(f"  - profit_per_unit: {updated['profit_per_unit']}")
    log(f"  - margin_pct: {updated['margin_pct']}")
    
    expected_profit = 10000 - 8000
    expected_margin = 20.0
    if abs(updated['profit_per_unit'] - expected_profit) > 0.01:
        error(f"Updated profit wrong: expected {expected_profit}, got {updated['profit_per_unit']}")
        return False
    if abs(updated['margin_pct'] - expected_margin) > 0.1:
        error(f"Updated margin wrong: expected {expected_margin}%, got {updated['margin_pct']}%")
        return False
    success(f"Product updated with new cost_price=8000, profit_per_unit={expected_profit}, margin_pct={expected_margin}%")
    
    # Test 1e: Check stock movements for the product
    log("Testing GET /api/admin/stock/movements for product...")
    resp = requests.get(f"{BASE_URL}/admin/stock/movements", params={"product_id": product_id}, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to get stock movements: {resp.status_code}")
        return False
    movements = resp.json()
    log(f"Found {len(movements)} stock movements")
    
    # Should have 'in' movement (initial stock 50) and 'adjust' movement (+10)
    if len(movements) < 2:
        error(f"Expected at least 2 movements (in + adjust), got {len(movements)}")
        return False
    
    # Check movements
    in_movement = None
    adjust_movement = None
    for m in movements:
        log(f"  - {m['movement_type']}: qty={m['qty']}, stock_before={m['stock_before']}, stock_after={m['stock_after']}, source={m['source']}")
        if m['movement_type'] == 'in' and m['stock_before'] == 0 and m['stock_after'] == 50:
            in_movement = m
        if m['movement_type'] == 'adjust' and m['stock_before'] == 50 and m['stock_after'] == 60:
            adjust_movement = m
    
    if not in_movement:
        error("Missing 'in' movement for initial stock (0 -> 50)")
        return False
    if not adjust_movement:
        error("Missing 'adjust' movement for stock update (50 -> 60)")
        return False
    
    success("Stock movements correctly logged: 'in' (stok awal 50) and 'adjust' (+10)")
    
    return product_id  # Return for use in other tests

def test_stock_endpoints(product_id: str):
    """Test 3: Stock endpoints"""
    log("\n=== TEST 3: Stock endpoints ===")
    
    # Test 3a: GET /api/admin/stock
    log("Testing GET /api/admin/stock...")
    resp = requests.get(f"{BASE_URL}/admin/stock", headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to get stock summary: {resp.status_code}")
        return False
    stock_data = resp.json()
    log(f"Stock summary:")
    log(f"  - total_products: {stock_data['total_products']}")
    log(f"  - total_units: {stock_data['total_units']}")
    log(f"  - total_stock_value: {stock_data['total_stock_value']}")
    log(f"  - out_of_stock: {stock_data['out_of_stock']}")
    log(f"  - low_stock: {stock_data['low_stock']}")
    
    # Find our test product
    test_product = None
    for item in stock_data['items']:
        if item['id'] == product_id:
            test_product = item
            break
    
    if not test_product:
        error(f"Test product {product_id} not found in stock summary")
        return False
    
    log(f"Test product in stock:")
    log(f"  - name: {test_product['name']}")
    log(f"  - stock: {test_product['stock']}")
    log(f"  - cost_price: {test_product['cost_price']}")
    log(f"  - stock_value: {test_product['stock_value']}")
    log(f"  - status: {test_product['status']}")
    
    # Verify stock_value = stock * cost_price
    expected_value = test_product['stock'] * test_product['cost_price']
    if abs(test_product['stock_value'] - expected_value) > 0.01:
        error(f"Stock value wrong: expected {expected_value}, got {test_product['stock_value']}")
        return False
    
    # Verify status (stock=60, min_order=5, so should be "aman")
    if test_product['status'] not in ['habis', 'menipis', 'aman']:
        error(f"Invalid status: {test_product['status']}")
        return False
    
    success(f"Stock summary correct: stock_value={test_product['stock_value']}, status={test_product['status']}")
    
    # Test 3b: POST /api/admin/stock/{id}/adjust - add stock (in)
    log("Testing POST /api/admin/stock/{id}/adjust with movement_type='in'...")
    resp = requests.post(f"{BASE_URL}/admin/stock/{product_id}/adjust", json={
        "movement_type": "in",
        "qty": 10,
        "note": "Tambah stok untuk testing"
    }, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to adjust stock (in): {resp.status_code} - {resp.text}")
        return False
    result = resp.json()
    log(f"Stock adjusted (in): stock={result['stock']}")
    if result['stock'] != 70:  # Was 60, added 10
        error(f"Stock after 'in' should be 70, got {result['stock']}")
        return False
    success("Stock 'in' adjustment successful: 60 -> 70")
    
    # Test 3c: POST /api/admin/stock/{id}/adjust - reduce stock (out)
    log("Testing POST /api/admin/stock/{id}/adjust with movement_type='out'...")
    resp = requests.post(f"{BASE_URL}/admin/stock/{product_id}/adjust", json={
        "movement_type": "out",
        "qty": 5,
        "note": "Kurangi stok untuk testing"
    }, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to adjust stock (out): {resp.status_code} - {resp.text}")
        return False
    result = resp.json()
    log(f"Stock adjusted (out): stock={result['stock']}")
    if result['stock'] != 65:  # Was 70, reduced 5
        error(f"Stock after 'out' should be 65, got {result['stock']}")
        return False
    success("Stock 'out' adjustment successful: 70 -> 65")
    
    # Test 3d: POST /api/admin/stock/{id}/adjust - out with qty > stock (should fail)
    log("Testing POST /api/admin/stock/{id}/adjust with movement_type='out' qty > stock (should fail)...")
    resp = requests.post(f"{BASE_URL}/admin/stock/{product_id}/adjust", json={
        "movement_type": "out",
        "qty": 999999
    }, headers=get_headers())
    if resp.status_code != 400:
        error(f"Expected 400 for out > stock, got {resp.status_code}")
        return False
    success("Stock 'out' with qty > stock correctly rejected with 400")
    
    # Test 3e: POST /api/admin/stock/{id}/adjust - adjust to 0
    log("Testing POST /api/admin/stock/{id}/adjust with movement_type='adjust' qty=0...")
    resp = requests.post(f"{BASE_URL}/admin/stock/{product_id}/adjust", json={
        "movement_type": "adjust",
        "qty": 0,
        "note": "Set stok ke 0 untuk testing"
    }, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to adjust stock to 0: {resp.status_code} - {resp.text}")
        return False
    result = resp.json()
    log(f"Stock adjusted to 0: stock={result['stock']}, status={result['status']}")
    if result['stock'] != 0:
        error(f"Stock should be 0, got {result['stock']}")
        return False
    if result['status'] != 'habis':
        error(f"Status should be 'habis', got {result['status']}")
        return False
    success("Stock 'adjust' to 0 successful, status='habis'")
    
    # Test 3f: POST /api/admin/stock/{id}/adjust - in with qty=0 (should fail)
    log("Testing POST /api/admin/stock/{id}/adjust with movement_type='in' qty=0 (should fail)...")
    resp = requests.post(f"{BASE_URL}/admin/stock/{product_id}/adjust", json={
        "movement_type": "in",
        "qty": 0
    }, headers=get_headers())
    if resp.status_code != 400:
        error(f"Expected 400 for in with qty=0, got {resp.status_code}")
        return False
    success("Stock 'in' with qty=0 correctly rejected with 400")
    
    # Restore stock for checkout test
    log("Restoring stock to 100 for checkout test...")
    resp = requests.post(f"{BASE_URL}/admin/stock/{product_id}/adjust", json={
        "movement_type": "adjust",
        "qty": 100
    }, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to restore stock: {resp.status_code}")
        return False
    success("Stock restored to 100")
    
    # Test 3g: Verify all movements are logged
    log("Verifying all stock movements are logged...")
    resp = requests.get(f"{BASE_URL}/admin/stock/movements", params={"product_id": product_id}, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to get stock movements: {resp.status_code}")
        return False
    movements = resp.json()
    log(f"Total movements: {len(movements)}")
    for m in movements:
        log(f"  - {m['movement_type']}: qty={m['qty']}, {m['stock_before']} -> {m['stock_after']}, source={m['source']}, created_by={m.get('created_by', 'N/A')}")
        if m.get('created_by') != 'admin':
            error(f"Movement created_by should be 'admin', got {m.get('created_by')}")
            return False
        if m['source'] != 'manual':
            error(f"Movement source should be 'manual', got {m['source']}")
            return False
    
    success("All stock movements correctly logged with created_by='admin' and source='manual'")
    return True

def test_checkout_flow(product_id: str):
    """Test 4: Checkout flow"""
    log("\n=== TEST 4: Checkout flow ===")
    
    # Get product details first
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    test_product = None
    for p in products:
        if p['id'] == product_id:
            test_product = p
            break
    
    if not test_product:
        error(f"Test product {product_id} not found")
        return None
    
    log(f"Using product: {test_product['name']}")
    log(f"  - min_order: {test_product['min_order']}")
    log(f"  - stock: {test_product['stock']}")
    log(f"  - cost_price: {test_product['cost_price']}")
    
    # Test 4a: POST /api/checkout
    log("Testing POST /api/checkout...")
    checkout_data = {
        "full_name": "Budi Santoso",
        "phone": "081234567890",
        "address": "Jl. Merdeka No. 123, Jakarta Pusat, DKI Jakarta 10110",
        "payment_method": "cod",
        "items": [
            {
                "product_id": product_id,
                "qty": test_product['min_order']
            }
        ]
    }
    resp = requests.post(f"{BASE_URL}/checkout", json=checkout_data)
    if resp.status_code != 200:
        error(f"Checkout failed: {resp.status_code} - {resp.text}")
        return None
    checkout_result = resp.json()
    order = checkout_result['order']
    order_id = order['id']
    order_number = order['order_number']
    
    log(f"Checkout successful!")
    log(f"  - order_number: {order_number}")
    log(f"  - order_id: {order_id}")
    log(f"  - total: {order['total']}")
    log(f"  - payment_method: {order['payment_method']}")
    log(f"  - order_status: {order['order_status']}")
    
    # Verify order_number starts with "SJ-"
    if not order_number.startswith("SJ-"):
        error(f"Order number should start with 'SJ-', got {order_number}")
        return None
    success(f"Order number starts with 'SJ-': {order_number}")
    
    # Verify order items have cost_price
    if not order['items']:
        error("Order has no items")
        return None
    
    order_item = order['items'][0]
    log(f"Order item:")
    log(f"  - product_name: {order_item['product_name']}")
    log(f"  - qty: {order_item['qty']}")
    log(f"  - price: {order_item['price']}")
    log(f"  - cost_price: {order_item.get('cost_price')}")
    
    if order_item.get('cost_price') is None:
        error("Order item missing cost_price")
        return None
    
    if abs(order_item['cost_price'] - test_product['cost_price']) > 0.01:
        error(f"Order item cost_price {order_item['cost_price']} != product cost_price {test_product['cost_price']}")
        return None
    
    success(f"Order item has correct cost_price snapshot: {order_item['cost_price']}")
    
    # Test 4b: Verify stock decreased
    log("Verifying stock decreased...")
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    updated_product = None
    for p in products:
        if p['id'] == product_id:
            updated_product = p
            break
    
    expected_stock = test_product['stock'] - test_product['min_order']
    if updated_product['stock'] != expected_stock:
        error(f"Stock should be {expected_stock}, got {updated_product['stock']}")
        return None
    success(f"Stock correctly decreased: {test_product['stock']} -> {updated_product['stock']}")
    
    # Test 4c: Verify stock movement logged
    log("Verifying stock movement logged...")
    resp = requests.get(f"{BASE_URL}/admin/stock/movements", params={"product_id": product_id}, headers=get_headers())
    movements = resp.json()
    
    # Find the order movement
    order_movement = None
    for m in movements:
        if m.get('source') == 'order' and m.get('reference') == order_number:
            order_movement = m
            break
    
    if not order_movement:
        error(f"Stock movement with source='order' and reference='{order_number}' not found")
        return None
    
    log(f"Order stock movement:")
    log(f"  - movement_type: {order_movement['movement_type']}")
    log(f"  - qty: {order_movement['qty']}")
    log(f"  - stock_before: {order_movement['stock_before']}")
    log(f"  - stock_after: {order_movement['stock_after']}")
    log(f"  - source: {order_movement['source']}")
    log(f"  - reference: {order_movement['reference']}")
    
    if order_movement['movement_type'] != 'out':
        error(f"Movement type should be 'out', got {order_movement['movement_type']}")
        return None
    if order_movement['source'] != 'order':
        error(f"Movement source should be 'order', got {order_movement['source']}")
        return None
    if order_movement['reference'] != order_number:
        error(f"Movement reference should be '{order_number}', got {order_movement['reference']}")
        return None
    
    success(f"Stock movement correctly logged with source='order', movement_type='out', reference='{order_number}'")
    
    return order_id

def test_dashboard(order_id: str):
    """Test 2: Dashboard finance"""
    log("\n=== TEST 2: Dashboard finance ===")
    
    # Test 2a: GET /api/admin/dashboard BEFORE marking order paid
    log("Testing GET /api/admin/dashboard BEFORE marking order paid/selesai...")
    resp = requests.get(f"{BASE_URL}/admin/dashboard", headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to get dashboard: {resp.status_code}")
        return False
    dashboard_before = resp.json()
    
    log(f"Dashboard BEFORE:")
    log(f"  - revenue_paid: {dashboard_before['revenue_paid']}")
    log(f"  - cost_paid: {dashboard_before['cost_paid']}")
    log(f"  - gross_profit: {dashboard_before['gross_profit']}")
    log(f"  - margin_pct: {dashboard_before['margin_pct']}")
    log(f"  - stock_value: {dashboard_before['stock_value']}")
    log(f"  - profit_by_product count: {len(dashboard_before['profit_by_product'])}")
    
    # Verify required fields exist
    required_fields = ['cost_paid', 'gross_profit', 'margin_pct', 'stock_value', 'profit_by_product']
    for field in required_fields:
        if field not in dashboard_before:
            error(f"Dashboard missing field: {field}")
            return False
    success("Dashboard has all required finance fields")
    
    # Test 2b: PATCH /api/admin/orders/{order_id}/status to mark as selesai
    log(f"Marking order {order_id} as 'selesai'...")
    resp = requests.patch(f"{BASE_URL}/admin/orders/{order_id}/status", json={
        "order_status": "selesai"
    }, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to update order status: {resp.status_code} - {resp.text}")
        return False
    success("Order marked as 'selesai'")
    
    # Test 2c: GET /api/admin/dashboard AFTER marking order paid
    log("Testing GET /api/admin/dashboard AFTER marking order selesai...")
    resp = requests.get(f"{BASE_URL}/admin/dashboard", headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to get dashboard: {resp.status_code}")
        return False
    dashboard_after = resp.json()
    
    log(f"Dashboard AFTER:")
    log(f"  - revenue_paid: {dashboard_after['revenue_paid']}")
    log(f"  - cost_paid: {dashboard_after['cost_paid']}")
    log(f"  - gross_profit: {dashboard_after['gross_profit']}")
    log(f"  - margin_pct: {dashboard_after['margin_pct']}")
    log(f"  - stock_value: {dashboard_after['stock_value']}")
    log(f"  - profit_by_product count: {len(dashboard_after['profit_by_product'])}")
    
    # Verify revenue increased
    if dashboard_after['revenue_paid'] <= dashboard_before['revenue_paid']:
        error(f"Revenue should increase after marking order selesai")
        return False
    success(f"Revenue increased: {dashboard_before['revenue_paid']} -> {dashboard_after['revenue_paid']}")
    
    # Verify gross_profit > 0
    if dashboard_after['gross_profit'] <= 0:
        error(f"Gross profit should be > 0, got {dashboard_after['gross_profit']}")
        return False
    success(f"Gross profit > 0: {dashboard_after['gross_profit']}")
    
    # Verify profit_by_product contains the product
    if not dashboard_after['profit_by_product']:
        error("profit_by_product is empty")
        return False
    
    log("Profit by product:")
    for p in dashboard_after['profit_by_product']:
        log(f"  - {p['product_name']}: qty={p['qty_sold']}, revenue={p['revenue']}, cost={p['cost']}, profit={p['profit']}, margin={p['margin_pct']}%")
    
    success("Dashboard correctly shows profit data after order marked selesai")
    return True

def test_order_edit_delete(product_id: str):
    """Test 6 & 7: Order edit and delete"""
    log("\n=== TEST 6 & 7: Order edit and delete ===")
    
    # Create a new order for testing
    log("Creating a new order for edit/delete testing...")
    checkout_data = {
        "full_name": "Siti Nurhaliza",
        "phone": "081298765432",
        "address": "Jl. Sudirman No. 456, Bandung, Jawa Barat 40123",
        "payment_method": "cod",
        "items": [
            {
                "product_id": product_id,
                "qty": 5
            }
        ]
    }
    resp = requests.post(f"{BASE_URL}/checkout", json=checkout_data)
    if resp.status_code != 200:
        error(f"Failed to create order: {resp.status_code} - {resp.text}")
        return False
    order = resp.json()['order']
    order_id = order['id']
    order_number = order['order_number']
    original_total = order['total']
    
    log(f"Created order: {order_number} (ID: {order_id})")
    log(f"  - customer_name: {order['customer_name']}")
    log(f"  - phone: {order['phone']}")
    log(f"  - address: {order['address'][:50]}...")
    log(f"  - total: {order['total']}")
    log(f"  - order_status: {order['order_status']}")
    
    # Get current stock
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    product = None
    for p in products:
        if p['id'] == product_id:
            product = p
            break
    stock_after_order = product['stock']
    log(f"Stock after order: {stock_after_order}")
    
    # Test 6a: PUT /api/admin/orders/{order_id} - update customer info and shipping_fee
    log("Testing PUT /api/admin/orders/{order_id} - update customer info and shipping_fee...")
    update_data = {
        "customer_name": "Siti Nurhaliza (Updated)",
        "phone": "081299999999",
        "address": "Jl. Gatot Subroto No. 789, Bandung, Jawa Barat 40124 (Alamat Baru Lengkap)",
        "notes": "Mohon dikirim pagi hari",
        "shipping_fee": 5000
    }
    resp = requests.put(f"{BASE_URL}/admin/orders/{order_id}", json=update_data, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to update order: {resp.status_code} - {resp.text}")
        return False
    updated_order = resp.json()
    
    log(f"Order updated:")
    log(f"  - customer_name: {updated_order['customer_name']}")
    log(f"  - phone: {updated_order['phone']}")
    log(f"  - address: {updated_order['address'][:50]}...")
    log(f"  - notes: {updated_order['notes']}")
    log(f"  - shipping_fee: {updated_order['shipping_fee']}")
    log(f"  - total: {updated_order['total']}")
    
    # Verify updates
    if updated_order['customer_name'] != update_data['customer_name']:
        error(f"Customer name not updated")
        return False
    if updated_order['phone'] != update_data['phone']:
        error(f"Phone not updated")
        return False
    if updated_order['address'] != update_data['address']:
        error(f"Address not updated")
        return False
    if updated_order['notes'] != update_data['notes']:
        error(f"Notes not updated")
        return False
    if updated_order['shipping_fee'] != update_data['shipping_fee']:
        error(f"Shipping fee not updated")
        return False
    
    # Verify total = subtotal + shipping_fee
    expected_total = updated_order['subtotal'] + update_data['shipping_fee']
    if abs(updated_order['total'] - expected_total) > 0.01:
        error(f"Total should be {expected_total}, got {updated_order['total']}")
        return False
    
    success(f"Order updated successfully, total recalculated: {updated_order['total']}")
    
    # Test 6b: PUT /api/admin/orders/{order_id} - change status to dibatalkan
    log("Testing PUT /api/admin/orders/{order_id} - change status to 'dibatalkan'...")
    resp = requests.put(f"{BASE_URL}/admin/orders/{order_id}", json={
        "order_status": "dibatalkan"
    }, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to cancel order: {resp.status_code} - {resp.text}")
        return False
    success("Order status changed to 'dibatalkan'")
    
    # Verify stock restored
    log("Verifying stock restored after cancellation...")
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    product = None
    for p in products:
        if p['id'] == product_id:
            product = p
            break
    stock_after_cancel = product['stock']
    log(f"Stock after cancel: {stock_after_cancel} (was {stock_after_order})")
    
    if stock_after_cancel != stock_after_order + 5:  # Should restore 5 units
        error(f"Stock should be restored to {stock_after_order + 5}, got {stock_after_cancel}")
        return False
    success(f"Stock restored after cancellation: {stock_after_order} -> {stock_after_cancel}")
    
    # Verify stock movement with source='cancel'
    log("Verifying stock movement with source='cancel'...")
    resp = requests.get(f"{BASE_URL}/admin/stock/movements", params={"product_id": product_id}, headers=get_headers())
    movements = resp.json()
    cancel_movement = None
    for m in movements:
        if m.get('source') == 'cancel' and m.get('reference') == order_number:
            cancel_movement = m
            break
    
    if not cancel_movement:
        error(f"Stock movement with source='cancel' not found")
        return False
    log(f"Cancel movement: {cancel_movement['movement_type']}, qty={cancel_movement['qty']}, {cancel_movement['stock_before']} -> {cancel_movement['stock_after']}")
    success("Stock movement with source='cancel' logged correctly")
    
    # Test 6c: PUT /api/admin/orders/{order_id} - change status back to 'diproses'
    log("Testing PUT /api/admin/orders/{order_id} - change status to 'diproses'...")
    resp = requests.put(f"{BASE_URL}/admin/orders/{order_id}", json={
        "order_status": "diproses"
    }, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to reactivate order: {resp.status_code} - {resp.text}")
        return False
    success("Order status changed to 'diproses'")
    
    # Verify stock reduced again
    log("Verifying stock reduced again after reactivation...")
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    product = None
    for p in products:
        if p['id'] == product_id:
            product = p
            break
    stock_after_reactivate = product['stock']
    log(f"Stock after reactivate: {stock_after_reactivate} (was {stock_after_cancel})")
    
    if stock_after_reactivate != stock_after_cancel - 5:  # Should reduce 5 units again
        error(f"Stock should be {stock_after_cancel - 5}, got {stock_after_reactivate}")
        return False
    success(f"Stock reduced again after reactivation: {stock_after_cancel} -> {stock_after_reactivate}")
    
    # Test 7: DELETE order
    log("\n=== TEST 7: Order delete ===")
    
    # Create another order for delete testing
    log("Creating another order for delete testing...")
    checkout_data = {
        "full_name": "Ahmad Dahlan",
        "phone": "081277777777",
        "address": "Jl. Asia Afrika No. 100, Bandung, Jawa Barat 40111",
        "payment_method": "cod",
        "items": [
            {
                "product_id": product_id,
                "qty": 5
            }
        ]
    }
    resp = requests.post(f"{BASE_URL}/checkout", json=checkout_data)
    if resp.status_code != 200:
        error(f"Failed to create order for delete test: {resp.status_code}")
        return False
    delete_order = resp.json()['order']
    delete_order_id = delete_order['id']
    delete_order_number = delete_order['order_number']
    log(f"Created order for delete: {delete_order_number} (ID: {delete_order_id})")
    
    # Get stock before delete
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    product = None
    for p in products:
        if p['id'] == product_id:
            product = p
            break
    stock_before_delete = product['stock']
    log(f"Stock before delete: {stock_before_delete}")
    
    # Test 7a: DELETE /api/admin/orders/{id}
    log(f"Testing DELETE /api/admin/orders/{delete_order_id}...")
    resp = requests.delete(f"{BASE_URL}/admin/orders/{delete_order_id}", headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to delete order: {resp.status_code} - {resp.text}")
        return False
    delete_result = resp.json()
    log(f"Delete result: {delete_result}")
    
    if not delete_result.get('ok'):
        error("Delete result should have 'ok': true")
        return False
    success(f"Order deleted successfully: {delete_order_number}")
    
    # Verify stock restored
    log("Verifying stock restored after delete...")
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    product = None
    for p in products:
        if p['id'] == product_id:
            product = p
            break
    stock_after_delete = product['stock']
    log(f"Stock after delete: {stock_after_delete} (was {stock_before_delete})")
    
    if stock_after_delete != stock_before_delete + 5:
        error(f"Stock should be restored to {stock_before_delete + 5}, got {stock_after_delete}")
        return False
    success(f"Stock restored after delete: {stock_before_delete} -> {stock_after_delete}")
    
    # Verify order is gone
    log("Verifying order is deleted (GET should return 404)...")
    resp = requests.get(f"{BASE_URL}/admin/orders/{delete_order_id}", headers=get_headers())
    if resp.status_code != 404:
        error(f"Expected 404 for deleted order, got {resp.status_code}")
        return False
    success("Deleted order returns 404")
    
    # Test 7b: Delete an already cancelled order (stock should NOT be restored twice)
    log("\nTesting delete of already cancelled order...")
    
    # Create another order
    checkout_data = {
        "full_name": "Kartini Wijaya",
        "phone": "081266666666",
        "address": "Jl. Diponegoro No. 200, Semarang, Jawa Tengah 50241",
        "payment_method": "cod",
        "items": [
            {
                "product_id": product_id,
                "qty": 5
            }
        ]
    }
    resp = requests.post(f"{BASE_URL}/checkout", json=checkout_data)
    if resp.status_code != 200:
        error(f"Failed to create order: {resp.status_code}")
        return False
    cancel_order = resp.json()['order']
    cancel_order_id = cancel_order['id']
    log(f"Created order: {cancel_order['order_number']}")
    
    # Get stock
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    product = None
    for p in products:
        if p['id'] == product_id:
            product = p
            break
    stock_before_cancel = product['stock']
    
    # Cancel the order first
    log("Cancelling order first...")
    resp = requests.put(f"{BASE_URL}/admin/orders/{cancel_order_id}", json={
        "order_status": "dibatalkan"
    }, headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to cancel order: {resp.status_code}")
        return False
    
    # Get stock after cancel
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    product = None
    for p in products:
        if p['id'] == product_id:
            product = p
            break
    stock_after_cancel = product['stock']
    log(f"Stock after cancel: {stock_after_cancel} (was {stock_before_cancel})")
    
    # Now delete the cancelled order
    log("Deleting the cancelled order...")
    resp = requests.delete(f"{BASE_URL}/admin/orders/{cancel_order_id}", headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to delete cancelled order: {resp.status_code}")
        return False
    
    # Get stock after delete
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    products = resp.json()
    product = None
    for p in products:
        if p['id'] == product_id:
            product = p
            break
    stock_after_delete_cancelled = product['stock']
    log(f"Stock after delete cancelled order: {stock_after_delete_cancelled} (was {stock_after_cancel})")
    
    # Stock should NOT change (already restored during cancel)
    if stock_after_delete_cancelled != stock_after_cancel:
        error(f"Stock should NOT change when deleting cancelled order, was {stock_after_cancel}, now {stock_after_delete_cancelled}")
        return False
    success("Stock NOT restored twice when deleting already cancelled order")
    
    return True

def test_auth():
    """Test 8: Auth - endpoints without token should return 401/403"""
    log("\n=== TEST 8: Auth ===")
    
    log("Testing /api/admin/stock without token (should return 401/403)...")
    resp = requests.get(f"{BASE_URL}/admin/stock")
    if resp.status_code not in [401, 403]:
        error(f"Expected 401/403 for unauthorized request, got {resp.status_code}")
        return False
    success(f"Unauthorized request correctly rejected with {resp.status_code}")
    
    return True

def cleanup_test_products():
    """Test 9: Cleanup - delete test products"""
    log("\n=== TEST 9: Cleanup ===")
    
    log("Getting all products to find test products...")
    resp = requests.get(f"{BASE_URL}/admin/products", headers=get_headers())
    if resp.status_code != 200:
        error(f"Failed to get products: {resp.status_code}")
        return False
    
    products = resp.json()
    test_products = [p for p in products if "Test" in p['name'] or "test" in p['name'].lower()]
    
    log(f"Found {len(test_products)} test products to delete")
    
    for p in test_products:
        log(f"Deleting product: {p['name']} (ID: {p['id']})")
        resp = requests.delete(f"{BASE_URL}/admin/products/{p['id']}", headers=get_headers())
        if resp.status_code != 200:
            error(f"Failed to delete product {p['id']}: {resp.status_code}")
            continue
        success(f"Deleted product: {p['name']}")
    
    log("Note: Orders referencing deleted products will have product_id=NULL (expected behavior)")
    return True

def test_sales_report_login():
    """Test 10: Sales Report - Login endpoints"""
    log("\n=== TEST 10: Sales Report - Login ===")
    
    # Test owner login
    log("Testing POST /api/admin/login with owner credentials...")
    resp = requests.post(f"{BASE_URL}/admin/login", json={
        "username": "owner",
        "password": "owner123"
    })
    if resp.status_code != 200:
        error(f"Owner login failed: {resp.status_code} - {resp.text}")
        return False
    owner_data = resp.json()
    if "access_token" not in owner_data:
        error("Owner login response missing access_token")
        return False
    success(f"Owner login successful, token: {owner_data['access_token'][:20]}...")
    
    # Test admin login
    log("Testing POST /api/admin/login with admin credentials...")
    resp = requests.post(f"{BASE_URL}/admin/login", json={
        "username": "admin",
        "password": "admin123"
    })
    if resp.status_code != 200:
        error(f"Admin login failed: {resp.status_code} - {resp.text}")
        return False
    admin_data = resp.json()
    if "access_token" not in admin_data:
        error("Admin login response missing access_token")
        return False
    success(f"Admin login successful, token: {admin_data['access_token'][:20]}...")
    
    return True

def test_sales_report_basic():
    """Test 11: Sales Report - Basic GET endpoint"""
    log("\n=== TEST 11: Sales Report - Basic GET ===")
    
    log("Testing GET /api/admin/reports/sales (no params) with OWNER token...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales", headers=get_owner_headers())
    if resp.status_code != 200:
        error(f"Sales report failed: {resp.status_code} - {resp.text}")
        return False
    
    data = resp.json()
    log(f"Sales report response keys: {list(data.keys())}")
    
    # Verify required fields
    required_fields = ["start_date", "end_date", "summary", "rows", "daily", "top_products"]
    for field in required_fields:
        if field not in data:
            error(f"Sales report missing field: {field}")
            return False
    success("Sales report has all required fields")
    
    # Verify summary fields
    summary = data["summary"]
    log(f"Summary: {json.dumps(summary, indent=2)}")
    summary_fields = ["gross_revenue", "total_cost", "net_profit", "margin_pct", "paid_orders", "items_sold", "avg_order_value"]
    for field in summary_fields:
        if field not in summary:
            error(f"Summary missing field: {field}")
            return False
    success("Summary has all required fields")
    
    # Verify net_profit == gross_revenue - total_cost
    expected_profit = summary["gross_revenue"] - summary["total_cost"]
    if abs(summary["net_profit"] - expected_profit) > 0.01:
        error(f"Net profit calculation wrong: expected {expected_profit}, got {summary['net_profit']}")
        return False
    success(f"Net profit calculation correct: {summary['gross_revenue']} - {summary['total_cost']} = {summary['net_profit']}")
    
    # Verify paid_orders == len(rows)
    if summary["paid_orders"] != len(data["rows"]):
        error(f"Paid orders count mismatch: summary says {summary['paid_orders']}, but rows has {len(data['rows'])}")
        return False
    success(f"Paid orders count matches: {summary['paid_orders']} orders")
    
    # Verify all rows have payment_status=='paid' OR (payment_method=='cod' and order_status=='selesai')
    log(f"Verifying {len(data['rows'])} rows...")
    for i, row in enumerate(data["rows"]):
        is_paid = row["payment_status"] == "paid"
        is_cod_selesai = row["payment_method"] == "cod" and row["order_status"] == "selesai"
        is_cancelled = row["order_status"] == "dibatalkan"
        
        if is_cancelled:
            error(f"Row {i+1} has order_status='dibatalkan', should not be in report")
            return False
        
        if not (is_paid or is_cod_selesai):
            error(f"Row {i+1} is neither paid nor COD selesai: payment_status={row['payment_status']}, payment_method={row['payment_method']}, order_status={row['order_status']}")
            return False
    
    success(f"All {len(data['rows'])} rows are valid (paid OR COD selesai, not cancelled)")
    
    log(f"Report period: {data['start_date']} to {data['end_date']}")
    log(f"Total revenue: {summary['gross_revenue']}")
    log(f"Total cost: {summary['total_cost']}")
    log(f"Net profit: {summary['net_profit']}")
    log(f"Margin: {summary['margin_pct']}%")
    
    return True

def test_sales_report_date_filters():
    """Test 12: Sales Report - Date filters"""
    log("\n=== TEST 12: Sales Report - Date filters ===")
    
    from datetime import datetime, timedelta
    
    # Test with 60 days ago
    today = datetime.now()
    start_60d = (today - timedelta(days=60)).strftime("%Y-%m-%d")
    end_today = today.strftime("%Y-%m-%d")
    
    log(f"Testing GET /api/admin/reports/sales?start={start_60d}&end={end_today}...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales", 
                       params={"start": start_60d, "end": end_today},
                       headers=get_owner_headers())
    if resp.status_code != 200:
        error(f"Sales report with date filter failed: {resp.status_code} - {resp.text}")
        return False
    
    data_60d = resp.json()
    log(f"60-day report: {len(data_60d['rows'])} orders, revenue={data_60d['summary']['gross_revenue']}")
    
    # Test with default (current month)
    resp = requests.get(f"{BASE_URL}/admin/reports/sales", headers=get_owner_headers())
    data_month = resp.json()
    log(f"Current month report: {len(data_month['rows'])} orders, revenue={data_month['summary']['gross_revenue']}")
    
    # 60-day should have >= current month orders (sample data has 8 in 60d, 6 in current month)
    if len(data_60d['rows']) < len(data_month['rows']):
        error(f"60-day report should have >= current month orders, got {len(data_60d['rows'])} vs {len(data_month['rows'])}")
        return False
    success(f"60-day report has more/equal orders than current month: {len(data_60d['rows'])} >= {len(data_month['rows'])}")
    
    # Test with empty period (2020)
    log("Testing GET /api/admin/reports/sales?start=2020-01-01&end=2020-01-31...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales",
                       params={"start": "2020-01-01", "end": "2020-01-31"},
                       headers=get_owner_headers())
    if resp.status_code != 200:
        error(f"Empty period report failed: {resp.status_code} - {resp.text}")
        return False
    
    data_empty = resp.json()
    log(f"Empty period report: {len(data_empty['rows'])} orders, revenue={data_empty['summary']['gross_revenue']}")
    
    if len(data_empty['rows']) != 0:
        error(f"Empty period should have 0 orders, got {len(data_empty['rows'])}")
        return False
    if data_empty['summary']['gross_revenue'] != 0:
        error(f"Empty period should have 0 revenue, got {data_empty['summary']['gross_revenue']}")
        return False
    success("Empty period (2020) returns 0 rows and zero summary")
    
    return True

def test_sales_report_rbac():
    """Test 13: Sales Report - RBAC (403 for admin, 401 without token)"""
    log("\n=== TEST 13: Sales Report - RBAC ===")
    
    # Test with ADMIN token (should be 403)
    log("Testing GET /api/admin/reports/sales with ADMIN (non-owner) token (should be 403)...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales", headers=get_headers())
    if resp.status_code != 403:
        error(f"Expected 403 for admin token, got {resp.status_code}")
        return False
    success("Admin (non-owner) token correctly rejected with 403")
    
    # Test without token (should be 401)
    log("Testing GET /api/admin/reports/sales without token (should be 401)...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales")
    if resp.status_code != 401:
        error(f"Expected 401 for no token, got {resp.status_code}")
        return False
    success("No token correctly rejected with 401")
    
    # Test export.xlsx with admin token (should be 403)
    log("Testing GET /api/admin/reports/sales/export.xlsx with ADMIN token (should be 403)...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales/export.xlsx", headers=get_headers())
    if resp.status_code != 403:
        error(f"Expected 403 for admin token on export.xlsx, got {resp.status_code}")
        return False
    success("Admin token on export.xlsx correctly rejected with 403")
    
    # Test export.pdf without token (should be 401)
    log("Testing GET /api/admin/reports/sales/export.pdf without token (should be 401)...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales/export.pdf")
    if resp.status_code != 401:
        error(f"Expected 401 for no token on export.pdf, got {resp.status_code}")
        return False
    success("No token on export.pdf correctly rejected with 401")
    
    return True

def test_sales_report_validation():
    """Test 14: Sales Report - Validation (start > end, invalid dates)"""
    log("\n=== TEST 14: Sales Report - Validation ===")
    
    # Test start > end (should be 400)
    log("Testing GET /api/admin/reports/sales?start=2024-12-31&end=2024-01-01 (start > end, should be 400)...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales",
                       params={"start": "2024-12-31", "end": "2024-01-01"},
                       headers=get_owner_headers())
    if resp.status_code != 400:
        error(f"Expected 400 for start > end, got {resp.status_code}")
        return False
    
    # Check for Indonesian error message
    error_data = resp.json()
    if "detail" in error_data:
        log(f"Error message: {error_data['detail']}")
        if "tanggal" not in error_data["detail"].lower():
            error("Error message should be in Indonesian (contain 'tanggal')")
            return False
    success("start > end correctly rejected with 400 and Indonesian message")
    
    # Test invalid date format (should be 400)
    log("Testing GET /api/admin/reports/sales?start=abc&end=2024-12-31 (invalid date, should be 400)...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales",
                       params={"start": "abc", "end": "2024-12-31"},
                       headers=get_owner_headers())
    if resp.status_code != 400:
        error(f"Expected 400 for invalid date, got {resp.status_code}")
        return False
    success("Invalid date format correctly rejected with 400")
    
    return True

def test_sales_report_export_xlsx():
    """Test 15: Sales Report - Export Excel"""
    log("\n=== TEST 15: Sales Report - Export Excel ===")
    
    from datetime import datetime, timedelta
    
    today = datetime.now()
    start = (today - timedelta(days=30)).strftime("%Y-%m-%d")
    end = today.strftime("%Y-%m-%d")
    
    log(f"Testing GET /api/admin/reports/sales/export.xlsx?start={start}&end={end}...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales/export.xlsx",
                       params={"start": start, "end": end},
                       headers=get_owner_headers(),
                       timeout=30)
    
    if resp.status_code != 200:
        error(f"Excel export failed: {resp.status_code} - {resp.text}")
        return False
    
    # Check content-type
    content_type = resp.headers.get("content-type", "")
    log(f"Content-Type: {content_type}")
    if "spreadsheetml" not in content_type:
        error(f"Expected spreadsheetml content-type, got {content_type}")
        return False
    success("Content-Type is correct (spreadsheetml)")
    
    # Check Content-Disposition
    content_disp = resp.headers.get("content-disposition", "")
    log(f"Content-Disposition: {content_disp}")
    if "attachment" not in content_disp:
        error(f"Expected attachment in Content-Disposition, got {content_disp}")
        return False
    if f"Laporan-Penjualan-Semoyo-Joyo_{start}_{end}.xlsx" not in content_disp:
        error(f"Filename should contain Laporan-Penjualan-Semoyo-Joyo_{start}_{end}.xlsx")
        return False
    success(f"Content-Disposition correct with filename")
    
    # Check body starts with PK (ZIP signature for xlsx)
    if not resp.content.startswith(b'PK'):
        error("Excel file should start with 'PK' (ZIP signature)")
        return False
    success("Excel file has valid ZIP signature (PK)")
    
    # Try to parse with openpyxl if available
    try:
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(resp.content))
        sheet_names = wb.sheetnames
        log(f"Excel sheet names: {sheet_names}")
        
        expected_sheets = ["Ringkasan", "Detail Pesanan", "Produk Terjual"]
        for sheet in expected_sheets:
            if sheet not in sheet_names:
                error(f"Missing sheet: {sheet}")
                return False
        success(f"Excel has all expected sheets: {expected_sheets}")
        
        # Check Detail Pesanan header row 4
        ws = wb["Detail Pesanan"]
        header_row = 4
        headers = [cell.value for cell in ws[header_row]]
        log(f"Detail Pesanan headers (row {header_row}): {headers}")
        
        expected_headers = ["No. Pesanan", "Nama Pembeli", "Total Penjualan", "Laba"]
        for h in expected_headers:
            if h not in headers:
                error(f"Missing header in Detail Pesanan: {h}")
                return False
        success(f"Detail Pesanan has all expected headers")
        
    except ImportError:
        log("openpyxl not available, skipping detailed Excel validation")
    except Exception as e:
        error(f"Failed to parse Excel: {e}")
        return False
    
    return True

def test_sales_report_export_pdf():
    """Test 16: Sales Report - Export PDF"""
    log("\n=== TEST 16: Sales Report - Export PDF ===")
    
    from datetime import datetime, timedelta
    
    today = datetime.now()
    start = (today - timedelta(days=30)).strftime("%Y-%m-%d")
    end = today.strftime("%Y-%m-%d")
    
    log(f"Testing GET /api/admin/reports/sales/export.pdf?start={start}&end={end}...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales/export.pdf",
                       params={"start": start, "end": end},
                       headers=get_owner_headers(),
                       timeout=30)
    
    if resp.status_code != 200:
        error(f"PDF export failed: {resp.status_code} - {resp.text}")
        return False
    
    # Check content-type
    content_type = resp.headers.get("content-type", "")
    log(f"Content-Type: {content_type}")
    if "application/pdf" not in content_type:
        error(f"Expected application/pdf content-type, got {content_type}")
        return False
    success("Content-Type is correct (application/pdf)")
    
    # Check Content-Disposition
    content_disp = resp.headers.get("content-disposition", "")
    log(f"Content-Disposition: {content_disp}")
    if ".pdf" not in content_disp:
        error(f"Filename should contain .pdf")
        return False
    success("Content-Disposition filename contains .pdf")
    
    # Check body starts with %PDF
    if not resp.content.startswith(b'%PDF'):
        error("PDF file should start with '%PDF'")
        return False
    success("PDF file has valid signature (%PDF)")
    
    # Test empty period (2020) - should still return valid PDF
    log("Testing PDF export with empty period (2020)...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales/export.pdf",
                       params={"start": "2020-01-01", "end": "2020-01-31"},
                       headers=get_owner_headers(),
                       timeout=30)
    
    if resp.status_code != 200:
        error(f"Empty period PDF export failed: {resp.status_code}")
        return False
    if not resp.content.startswith(b'%PDF'):
        error("Empty period PDF should still be valid")
        return False
    success("Empty period (2020) returns valid PDF with 200")
    
    return True

def test_sales_report_audit_logs():
    """Test 17: Sales Report - Audit logs after exports"""
    log("\n=== TEST 17: Sales Report - Audit logs ===")
    
    from datetime import datetime, timedelta
    
    today = datetime.now()
    start = (today - timedelta(days=7)).strftime("%Y-%m-%d")
    end = today.strftime("%Y-%m-%d")
    
    # Do an export first
    log("Exporting Excel to generate audit log...")
    resp = requests.get(f"{BASE_URL}/admin/reports/sales/export.xlsx",
                       params={"start": start, "end": end},
                       headers=get_owner_headers(),
                       timeout=30)
    if resp.status_code != 200:
        error(f"Excel export failed: {resp.status_code}")
        return False
    
    # Check audit logs
    log("Testing GET /api/admin/audit-logs?action=export...")
    resp = requests.get(f"{BASE_URL}/admin/audit-logs",
                       params={"action": "export"},
                       headers=get_owner_headers())
    
    if resp.status_code != 200:
        error(f"Audit logs failed: {resp.status_code} - {resp.text}")
        return False
    
    logs = resp.json()
    log(f"Found {len(logs)} export audit logs")
    
    # Find report export logs
    report_logs = [l for l in logs if l.get("entity_type") == "report"]
    log(f"Found {len(report_logs)} report export logs")
    
    if len(report_logs) == 0:
        error("No report export logs found with entity_type='report'")
        return False
    
    # Check the most recent one
    latest = report_logs[0]
    log(f"Latest report export log:")
    log(f"  - action: {latest.get('action')}")
    log(f"  - entity_type: {latest.get('entity_type')}")
    log(f"  - entity_label: {latest.get('entity_label')}")
    log(f"  - description: {latest.get('description')}")
    
    if latest.get("action") != "export":
        error(f"Expected action='export', got {latest.get('action')}")
        return False
    if latest.get("entity_type") != "report":
        error(f"Expected entity_type='report', got {latest.get('entity_type')}")
        return False
    
    success("Audit log correctly recorded report export with entity_type='report'")
    
    return True

def main():
    """Main test runner"""
    print("\n" + "="*80)
    print("BACKEND API TESTING - Semoyo Joyo B2B App - Sales Report Feature")
    print("="*80 + "\n")
    
    # Login as both admin and owner
    if not admin_login():
        error("Failed to login as admin. Aborting tests.")
        return
    
    if not owner_login():
        error("Failed to login as owner. Aborting tests.")
        return
    
    # Test 10: Sales Report - Login
    if not test_sales_report_login():
        error("Test 10 failed. Continuing...")
    
    # Test 11: Sales Report - Basic GET
    if not test_sales_report_basic():
        error("Test 11 failed. Aborting remaining sales report tests.")
        return
    
    # Test 12: Sales Report - Date filters
    if not test_sales_report_date_filters():
        error("Test 12 failed. Continuing...")
    
    # Test 13: Sales Report - RBAC
    if not test_sales_report_rbac():
        error("Test 13 failed. Continuing...")
    
    # Test 14: Sales Report - Validation
    if not test_sales_report_validation():
        error("Test 14 failed. Continuing...")
    
    # Test 15: Sales Report - Export Excel
    if not test_sales_report_export_xlsx():
        error("Test 15 failed. Continuing...")
    
    # Test 16: Sales Report - Export PDF
    if not test_sales_report_export_pdf():
        error("Test 16 failed. Continuing...")
    
    # Test 17: Sales Report - Audit logs
    if not test_sales_report_audit_logs():
        error("Test 17 failed. Continuing...")
    
    print("\n" + "="*80)
    print("ALL SALES REPORT TESTS COMPLETED")
    print("="*80 + "\n")

if __name__ == "__main__":
    main()
