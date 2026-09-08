"""Laporan Penjualan (khusus Owner): ringkasan periode, detail pesanan, export Excel (.xlsx) & PDF.

Definisi "terjual": pesanan LUNAS (payment_status=paid) ATAU COD yang sudah SELESAI, dan tidak dibatalkan
(sama dengan perhitungan laba/rugi di Dashboard). Filter tanggal memakai tanggal pesanan (WIB, UTC+7).
"""
from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from audit import log_action
from auth import get_owner_user
from database import get_db
from models import Order, OrderItem, Product, User

router = APIRouter(prefix="/admin/reports", tags=["reports"])

WIB = timezone(timedelta(hours=7))
BRAND = "Semoyo Joyo"
TAGLINE = "Solusi Belanja Terpercaya"
BRAND_BLUE = "#0B4EA2"
BRAND_YELLOW = "#F5C400"
LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "logo.png"

# Pesanan yang dihitung terjual: lunas ATAU COD selesai; tidak dibatalkan
SOLD_FILTER = ((Order.payment_status == "paid") | ((Order.payment_method == "cod") & (Order.order_status == "selesai"))) & (Order.order_status != "dibatalkan")

PAYMENT_LABEL = {"cod": "COD", "bank_transfer": "Transfer Bank", "qris": "QRIS", "ewallet": "E-Wallet"}
STATUS_LABEL = {"baru": "Baru", "diproses": "Diproses", "dikirim": "Dikirim", "selesai": "Selesai", "dibatalkan": "Dibatalkan"}
MONTHS_ID = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]


# ---------------- Schemas ----------------
class ReportRow(BaseModel):
    order_id: str
    order_number: str
    created_at: datetime
    paid_at: Optional[datetime] = None
    customer_name: str
    phone: str
    payment_method: str
    payment_channel: Optional[str] = None
    payment_status: str
    order_status: str
    items_count: int
    subtotal: float
    shipping_fee: float
    total: float
    cost: float
    profit: float
    margin_pct: float


class ReportSummary(BaseModel):
    gross_revenue: float
    total_cost: float
    net_profit: float
    margin_pct: float
    paid_orders: int
    items_sold: int
    avg_order_value: float


class DailyPoint(BaseModel):
    date: date
    orders: int
    revenue: float
    cost: float
    profit: float


class ProductPoint(BaseModel):
    product_name: str
    qty: int
    revenue: float
    cost: float
    profit: float


class SalesReportOut(BaseModel):
    start_date: date
    end_date: date
    generated_at: datetime
    summary: ReportSummary
    daily: list[DailyPoint]
    top_products: list[ProductPoint]
    rows: list[ReportRow]


# ---------------- Helpers ----------------
def _parse_range(start: Optional[str], end: Optional[str]) -> tuple[date, date]:
    today = datetime.now(WIB).date()
    try:
        s = date.fromisoformat(start) if start else today.replace(day=1)
        e = date.fromisoformat(end) if end else today
    except ValueError:
        raise HTTPException(400, "Format tanggal harus YYYY-MM-DD")
    if s > e:
        raise HTTPException(400, "Tanggal mulai tidak boleh setelah tanggal akhir")
    if (e - s).days > 366 * 3:
        raise HTTPException(400, "Rentang laporan maksimal 3 tahun")
    return s, e


def _bounds_utc(s: date, e: date) -> tuple[datetime, datetime]:
    start_dt = datetime.combine(s, time.min, tzinfo=WIB).astimezone(timezone.utc)
    end_dt = datetime.combine(e + timedelta(days=1), time.min, tzinfo=WIB).astimezone(timezone.utc)
    return start_dt, end_dt


def fmt_date_id(d: date) -> str:
    return f"{d.day} {MONTHS_ID[d.month]} {d.year}"


def fmt_dt_id(dt: Optional[datetime]) -> str:
    if not dt:
        return "-"
    local = dt.astimezone(WIB)
    return local.strftime("%d/%m/%Y %H:%M")


def rp(n: float) -> str:
    sign = "-" if n < 0 else ""
    return f"{sign}Rp {abs(n):,.0f}".replace(",", ".")


async def build_report(db: AsyncSession, s: date, e: date) -> SalesReportOut:
    start_dt, end_dt = _bounds_utc(s, e)
    cost_expr = func.coalesce(OrderItem.cost_price, Product.cost_price, 0) * OrderItem.qty
    agg = (
        select(OrderItem.order_id, func.sum(cost_expr).label("cost"), func.sum(OrderItem.qty).label("qty"))
        .outerjoin(Product, Product.id == OrderItem.product_id)
        .group_by(OrderItem.order_id)
        .subquery()
    )
    stmt = (
        select(Order, func.coalesce(agg.c.cost, 0), func.coalesce(agg.c.qty, 0))
        .outerjoin(agg, agg.c.order_id == Order.id)
        .where(SOLD_FILTER, Order.created_at >= start_dt, Order.created_at < end_dt)
        .order_by(Order.created_at.desc())
    )
    res = (await db.execute(stmt)).all()

    rows: list[ReportRow] = []
    daily: dict[date, DailyPoint] = {}
    total_rev = total_cost = 0.0
    items_sold = 0
    for o, cost, qty in res:
        rev = float(o.total or 0)
        cost_f = float(cost or 0)
        profit = rev - cost_f
        total_rev += rev
        total_cost += cost_f
        items_sold += int(qty or 0)
        rows.append(ReportRow(
            order_id=o.id, order_number=o.order_number, created_at=o.created_at, paid_at=o.paid_at, customer_name=o.customer_name, phone=o.phone,
            payment_method=o.payment_method, payment_channel=o.payment_channel, payment_status=o.payment_status, order_status=o.order_status,
            items_count=int(qty or 0), subtotal=float(o.subtotal or 0), shipping_fee=float(o.shipping_fee or 0), total=rev, cost=cost_f, profit=profit,
            margin_pct=round(profit / rev * 100, 1) if rev > 0 else 0.0,
        ))
        d = o.created_at.astimezone(WIB).date()
        pt = daily.setdefault(d, DailyPoint(date=d, orders=0, revenue=0.0, cost=0.0, profit=0.0))
        pt.orders += 1
        pt.revenue += rev
        pt.cost += cost_f
        pt.profit += profit

    # Produk terlaris pada periode
    prod_stmt = (
        select(func.max(OrderItem.product_name), func.sum(OrderItem.qty), func.sum(OrderItem.subtotal), func.sum(cost_expr))
        .join(Order, Order.id == OrderItem.order_id)
        .outerjoin(Product, Product.id == OrderItem.product_id)
        .where(SOLD_FILTER, Order.created_at >= start_dt, Order.created_at < end_dt)
        .group_by(func.coalesce(OrderItem.product_id, OrderItem.product_name))
    )
    top: list[ProductPoint] = []
    for name, q, rev, cost in (await db.execute(prod_stmt)).all():
        rev_f, cost_f = float(rev or 0), float(cost or 0)
        top.append(ProductPoint(product_name=name, qty=int(q or 0), revenue=rev_f, cost=cost_f, profit=rev_f - cost_f))
    top.sort(key=lambda r: r.revenue, reverse=True)

    n = len(rows)
    net = total_rev - total_cost
    summary = ReportSummary(
        gross_revenue=total_rev, total_cost=total_cost, net_profit=net, margin_pct=round(net / total_rev * 100, 1) if total_rev > 0 else 0.0,
        paid_orders=n, items_sold=items_sold, avg_order_value=round(total_rev / n, 2) if n else 0.0,
    )
    return SalesReportOut(start_date=s, end_date=e, generated_at=datetime.now(timezone.utc), summary=summary,
                          daily=sorted(daily.values(), key=lambda p: p.date), top_products=top[:15], rows=rows)


def _filename(s: date, e: date, ext: str) -> str:
    return f"Laporan-Penjualan-Semoyo-Joyo_{s.isoformat()}_{e.isoformat()}.{ext}"


# ---------------- Endpoints ----------------
@router.get("/sales", response_model=SalesReportOut)
async def sales_report(start: Optional[str] = Query(default=None), end: Optional[str] = Query(default=None),
                       _: User = Depends(get_owner_user), db: AsyncSession = Depends(get_db)):
    s, e = _parse_range(start, end)
    return await build_report(db, s, e)


@router.get("/sales/export.xlsx")
async def sales_export_xlsx(start: Optional[str] = Query(default=None), end: Optional[str] = Query(default=None),
                            owner: User = Depends(get_owner_user), db: AsyncSession = Depends(get_db)):
    s, e = _parse_range(start, end)
    report = await build_report(db, s, e)
    data = render_xlsx(report)
    log_action(db, owner, "export", "report", f"Mengunduh Laporan Penjualan (Excel) periode {s.isoformat()} s/d {e.isoformat()}",
               entity_label="Laporan Penjualan", meta={"format": "xlsx", "start": s.isoformat(), "end": e.isoformat(), "orders": report.summary.paid_orders})
    await db.commit()
    fname = _filename(s, e, "xlsx")
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/sales/export.pdf")
async def sales_export_pdf(start: Optional[str] = Query(default=None), end: Optional[str] = Query(default=None),
                           owner: User = Depends(get_owner_user), db: AsyncSession = Depends(get_db)):
    s, e = _parse_range(start, end)
    report = await build_report(db, s, e)
    data = render_pdf(report, owner.full_name)
    log_action(db, owner, "export", "report", f"Mengunduh Laporan Penjualan (PDF) periode {s.isoformat()} s/d {e.isoformat()}",
               entity_label="Laporan Penjualan", meta={"format": "pdf", "start": s.isoformat(), "end": e.isoformat(), "orders": report.summary.paid_orders})
    await db.commit()
    fname = _filename(s, e, "pdf")
    return StreamingResponse(io.BytesIO(data), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ---------------- Excel renderer ----------------
def render_xlsx(r: SalesReportOut) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    RP_FMT = '"Rp" #,##0;[Red]"-Rp" #,##0'
    blue = BRAND_BLUE.lstrip("#")
    yellow = BRAND_YELLOW.lstrip("#")
    head_fill = PatternFill("solid", fgColor=blue)
    head_font = Font(bold=True, color="FFFFFF")
    total_fill = PatternFill("solid", fgColor="FFF4C2")
    thin = Side(style="thin", color="D9DEE7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    periode = f"{fmt_date_id(r.start_date)} - {fmt_date_id(r.end_date)}"
    generated = fmt_dt_id(r.generated_at) + " WIB"

    wb = Workbook()

    # ---- Sheet 1: Ringkasan ----
    ws = wb.active
    ws.title = "Ringkasan"
    ws.sheet_view.showGridLines = False
    if LOGO_PATH.exists():
        img = XLImage(str(LOGO_PATH))
        ratio = img.height / img.width if img.width else 0.25
        img.width, img.height = 220, int(220 * ratio)
        ws.add_image(img, "B2")
    ws["B6"] = "LAPORAN PENJUALAN"
    ws["B6"].font = Font(bold=True, size=16, color=blue)
    ws["B7"] = f"{BRAND} - {TAGLINE}"
    ws["B7"].font = Font(italic=True, color="6B7280")
    ws["B9"], ws["C9"] = "Periode", periode
    ws["B10"], ws["C10"] = "Dibuat pada", generated
    ws["B11"], ws["C11"] = "Dasar perhitungan", "Pesanan lunas / COD selesai (tidak dibatalkan), berdasarkan tanggal pesanan"
    for c in ("B9", "B10", "B11"):
        ws[c].font = Font(bold=True)

    ws["B13"], ws["C13"] = "Indikator", "Nilai"
    for c in ("B13", "C13"):
        ws[c].fill, ws[c].font, ws[c].border = head_fill, head_font, border
    sm = r.summary
    metrics = [
        ("Total Pendapatan Kotor", sm.gross_revenue, RP_FMT),
        ("Total Modal (HPP)", sm.total_cost, RP_FMT),
        ("Total Laba Bersih", sm.net_profit, RP_FMT),
        ("Margin Laba", sm.margin_pct / 100, "0.0%"),
        ("Jumlah Pesanan Lunas", sm.paid_orders, "#,##0"),
        ("Total Item Terjual", sm.items_sold, "#,##0"),
        ("Rata-rata Nilai Pesanan", sm.avg_order_value, RP_FMT),
    ]
    for i, (label, val, fmt) in enumerate(metrics, start=14):
        ws.cell(row=i, column=2, value=label).border = border
        c = ws.cell(row=i, column=3, value=val)
        c.number_format, c.border, c.alignment = fmt, border, Alignment(horizontal="right")
        if label == "Total Laba Bersih":
            c.font = Font(bold=True, color="047857" if val >= 0 else "B91C1C")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 70

    # ---- Sheet 2: Detail Pesanan ----
    wd = wb.create_sheet("Detail Pesanan")
    headers = ["No", "No. Pesanan", "Tanggal Pesanan", "Tanggal Lunas", "Nama Pembeli", "No. Telp/WA", "Metode Bayar", "Status Pesanan", "Jumlah Item",
               "Subtotal", "Ongkir", "Total Penjualan", "Modal (HPP)", "Laba", "Margin %"]
    wd["A1"] = f"Detail Pesanan - {BRAND}"
    wd["A1"].font = Font(bold=True, size=13, color=blue)
    wd["A2"] = f"Periode: {periode}  |  Dibuat: {generated}"
    wd["A2"].font = Font(color="6B7280")
    HR = 4
    for col, h in enumerate(headers, start=1):
        c = wd.cell(row=HR, column=col, value=h)
        c.fill, c.font, c.border = head_fill, head_font, border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wd.row_dimensions[HR].height = 30
    for i, row in enumerate(r.rows, start=1):
        rr = HR + i
        values = [i, row.order_number, row.created_at.astimezone(WIB).replace(tzinfo=None), row.paid_at.astimezone(WIB).replace(tzinfo=None) if row.paid_at else "-",
                  row.customer_name, row.phone, PAYMENT_LABEL.get(row.payment_method, row.payment_method) + (f" ({row.payment_channel.upper()})" if row.payment_channel else ""),
                  STATUS_LABEL.get(row.order_status, row.order_status), row.items_count, row.subtotal, row.shipping_fee, row.total, row.cost, row.profit, row.margin_pct / 100]
        for col, v in enumerate(values, start=1):
            c = wd.cell(row=rr, column=col, value=v)
            c.border = border
            if col in (3, 4) and not isinstance(v, str):
                c.number_format = "dd/mm/yyyy hh:mm"
            elif col in (10, 11, 12, 13, 14):
                c.number_format = RP_FMT
            elif col == 15:
                c.number_format = "0.0%"
            elif col in (1, 9):
                c.alignment = Alignment(horizontal="center")
            if col == 14:
                c.font = Font(bold=True, color="047857" if (v or 0) >= 0 else "B91C1C")
    # Totals row
    tr = HR + len(r.rows) + 1
    wd.cell(row=tr, column=2, value="TOTAL").font = Font(bold=True)
    wd.cell(row=tr, column=9, value=r.summary.items_sold)
    first, last = HR + 1, HR + len(r.rows)
    for col in (10, 11, 12, 13, 14):
        L = get_column_letter(col)
        c = wd.cell(row=tr, column=col, value=f"=SUM({L}{first}:{L}{last})" if r.rows else 0)
        c.number_format = RP_FMT
        c.font = Font(bold=True)
    mc = wd.cell(row=tr, column=15, value=r.summary.margin_pct / 100)
    mc.number_format, mc.font = "0.0%", Font(bold=True)
    for col in range(1, len(headers) + 1):
        c = wd.cell(row=tr, column=col)
        c.fill, c.border = total_fill, border
    widths = [5, 20, 18, 18, 26, 16, 20, 14, 10, 16, 14, 18, 16, 16, 10]
    for col, w in enumerate(widths, start=1):
        wd.column_dimensions[get_column_letter(col)].width = w
    wd.freeze_panes = f"C{HR + 1}"
    if r.rows:
        wd.auto_filter.ref = f"A{HR}:{get_column_letter(len(headers))}{last}"

    # ---- Sheet 3: Produk Terjual ----
    wp = wb.create_sheet("Produk Terjual")
    wp["A1"] = f"Produk Terjual - {BRAND}"
    wp["A1"].font = Font(bold=True, size=13, color=blue)
    wp["A2"] = f"Periode: {periode}"
    wp["A2"].font = Font(color="6B7280")
    ph = ["No", "Nama Produk", "Qty Terjual", "Omzet", "Modal (HPP)", "Laba", "Margin %"]
    for col, h in enumerate(ph, start=1):
        c = wp.cell(row=4, column=col, value=h)
        c.fill, c.font, c.border, c.alignment = head_fill, head_font, border, Alignment(horizontal="center")
    for i, p in enumerate(r.top_products, start=1):
        vals = [i, p.product_name, p.qty, p.revenue, p.cost, p.profit, (p.profit / p.revenue) if p.revenue else 0]
        for col, v in enumerate(vals, start=1):
            c = wp.cell(row=4 + i, column=col, value=v)
            c.border = border
            if col in (4, 5, 6):
                c.number_format = RP_FMT
            elif col == 7:
                c.number_format = "0.0%"
    for col, w in enumerate([5, 36, 12, 18, 18, 18, 10], start=1):
        wp.column_dimensions[get_column_letter(col)].width = w
    wp.freeze_panes = "A5"

    # Fill yellow accent on summary header row
    ws["B6"].fill = PatternFill("solid", fgColor="FFFFFF")
    ws.cell(row=5, column=2).fill = PatternFill("solid", fgColor=yellow)
    ws.row_dimensions[5].height = 4

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- PDF renderer ----------------
def render_pdf(r: SalesReportOut, owner_name: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    blue = colors.HexColor(BRAND_BLUE)
    yellow = colors.HexColor(BRAND_YELLOW)
    grey = colors.HexColor("#6B7280")
    line = colors.HexColor("#D9DEE7")
    zebra = colors.HexColor("#F4F7FB")
    green = colors.HexColor("#047857")
    red = colors.HexColor("#B91C1C")

    styles = getSampleStyleSheet()
    h_title = ParagraphStyle("t", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, textColor=blue, alignment=TA_RIGHT, spaceAfter=0, leading=22)
    h_sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, textColor=grey, alignment=TA_RIGHT, leading=12)
    sec = ParagraphStyle("sec", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=11, textColor=blue, spaceBefore=6, spaceAfter=4)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7.5, leading=9.5)
    small_r = ParagraphStyle("small_r", parent=small, alignment=TA_RIGHT)
    kpi_label = ParagraphStyle("kl", parent=styles["Normal"], fontSize=7.5, textColor=grey, leading=10)
    kpi_val = ParagraphStyle("kv", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#111827"))
    note = ParagraphStyle("note", parent=styles["Normal"], fontSize=7.5, textColor=grey, leading=10)

    pagesize = landscape(A4)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title=f"Laporan Penjualan {BRAND}", author=BRAND, subject="Laporan Penjualan")
    W = pagesize[0] - doc.leftMargin - doc.rightMargin
    periode = f"{fmt_date_id(r.start_date)} - {fmt_date_id(r.end_date)}"
    generated = fmt_dt_id(r.generated_at) + " WIB"

    story = []
    # ---- Kop / letterhead ----
    logo_cell = ""
    if LOGO_PATH.exists():
        from PIL import Image as PILImage
        with PILImage.open(LOGO_PATH) as im:
            ratio = im.height / im.width
        lw = 62 * mm
        logo_cell = Image(str(LOGO_PATH), width=lw, height=lw * ratio)
    right = [Paragraph("LAPORAN PENJUALAN", h_title), Paragraph(f"Periode: <b>{periode}</b>", h_sub), Paragraph(f"Dibuat: {generated} &nbsp;|&nbsp; Oleh: {owner_name} (Owner)", h_sub)]
    kop = Table([[logo_cell, right]], colWidths=[W * 0.42, W * 0.58])
    kop.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                             ("LINEBELOW", (0, 0), (-1, 0), 2.2, blue), ("BOTTOMPADDING", (0, 0), (-1, 0), 6)]))
    story.append(kop)
    accent = Table([[""]], colWidths=[W], rowHeights=[2.2])
    accent.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), yellow), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    story += [accent, Spacer(1, 6 * mm)]

    # ---- Ringkasan KPI ----
    sm = r.summary
    profit_color = green if sm.net_profit >= 0 else red
    kpis = [
        ("Total Pendapatan Kotor", rp(sm.gross_revenue), None),
        ("Total Modal (HPP)", rp(sm.total_cost), None),
        ("Total Laba Bersih", rp(sm.net_profit), profit_color),
        ("Margin Laba", f"{sm.margin_pct:.1f}%", None),
        ("Jumlah Pesanan Lunas", f"{sm.paid_orders:,}".replace(",", "."), None),
    ]
    kpi_cells = []
    for label, val, col in kpis:
        vs = ParagraphStyle("kv2", parent=kpi_val, textColor=col) if col else kpi_val
        kpi_cells.append([Paragraph(label, kpi_label), Paragraph(val, vs)])
    kpi_tbl = Table([[Table([[cell] for cell in c], colWidths=[W / 5 - 6 * mm]) for c in kpi_cells]], colWidths=[W / 5] * 5)
    kpi_tbl.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.6, line), ("INNERGRID", (0, 0), (-1, -1), 0.6, line), ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAFBFD")),
        ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#ECFDF5") if sm.net_profit >= 0 else colors.HexColor("#FEF2F2")),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story += [Paragraph("Ringkasan Periode", sec), kpi_tbl, Spacer(1, 2 * mm),
              Paragraph(f"Dasar perhitungan: pesanan <b>lunas</b> atau <b>COD selesai</b> (tidak dibatalkan) berdasarkan tanggal pesanan. "
                        f"Item terjual: {sm.items_sold:,} &nbsp;|&nbsp; Rata-rata nilai pesanan: {rp(sm.avg_order_value)}".replace(",", "."), note),
              Spacer(1, 5 * mm)]

    # ---- Detail table ----
    story.append(Paragraph(f"Detail Pesanan ({len(r.rows)} pesanan)", sec))
    head = ["No", "No. Pesanan", "Tanggal", "Nama Pembeli", "Metode Bayar", "Status", "Item", "Total Penjualan", "Modal (HPP)", "Laba", "Margin"]
    data = [head]
    for i, row in enumerate(r.rows, start=1):
        method = PAYMENT_LABEL.get(row.payment_method, row.payment_method) + (f" ({row.payment_channel.upper()})" if row.payment_channel else "")
        data.append([str(i), row.order_number, fmt_dt_id(row.created_at), Paragraph(row.customer_name, small), method, STATUS_LABEL.get(row.order_status, row.order_status),
                     str(row.items_count), rp(row.total), rp(row.cost), rp(row.profit), f"{row.margin_pct:.1f}%"])
    if not r.rows:
        data.append(["", Paragraph("Tidak ada pesanan lunas pada periode ini.", small)] + [""] * 9)
    data.append(["", "TOTAL", "", "", "", "", str(sm.items_sold), rp(sm.gross_revenue), rp(sm.total_cost), rp(sm.net_profit), f"{sm.margin_pct:.1f}%"])
    cw = [8, 30, 26, 48, 32, 18, 10, 30, 28, 28, 14]
    scale = W / sum(cw)
    tbl = Table(data, colWidths=[c * scale for c in cw], repeatRows=1)
    ts = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), blue), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("ALIGN", (0, 0), (-1, 0), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 1), (0, -1), "CENTER"), ("ALIGN", (6, 1), (6, -1), "CENTER"), ("ALIGN", (7, 1), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, line), ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, zebra]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FFF4C2")), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
    ])
    for i, row in enumerate(r.rows, start=1):
        ts.add("TEXTCOLOR", (9, i), (9, i), green if row.profit >= 0 else red)
    ts.add("TEXTCOLOR", (9, -1), (9, -1), green if sm.net_profit >= 0 else red)
    tbl.setStyle(ts)
    story.append(tbl)

    # ---- Produk terjual ----
    if r.top_products:
        story += [Spacer(1, 6 * mm), Paragraph("Produk Terjual (Top 15 berdasarkan omzet)", sec)]
        pdata = [["No", "Nama Produk", "Qty", "Omzet", "Modal (HPP)", "Laba", "Margin"]]
        for i, p in enumerate(r.top_products, start=1):
            m = (p.profit / p.revenue * 100) if p.revenue else 0
            pdata.append([str(i), Paragraph(p.product_name, small), str(p.qty), rp(p.revenue), rp(p.cost), rp(p.profit), f"{m:.1f}%"])
        pcw = [8, 80, 14, 30, 30, 30, 14]
        pscale = W / sum(pcw)
        ptbl = Table(pdata, colWidths=[c * pscale for c in pcw], repeatRows=1)
        ptbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), blue), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("ALIGN", (0, 0), (-1, 0), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 1), (0, -1), "CENTER"), ("ALIGN", (2, 1), (-1, -1), "RIGHT"), ("GRID", (0, 0), (-1, -1), 0.4, line),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, zebra]), ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ]))
        story.append(ptbl)

    # ---- Tanda tangan ----
    story += [Spacer(1, 10 * mm)]
    sign = Table([["", Paragraph(f"Dicetak {generated}<br/><br/><br/><br/>________________________<br/><b>{owner_name}</b><br/>Owner {BRAND}", small_r)]], colWidths=[W * 0.7, W * 0.3])
    sign.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(sign)

    def on_page(canvas, d):
        canvas.saveState()
        canvas.setStrokeColor(line)
        canvas.setLineWidth(0.5)
        canvas.line(d.leftMargin, 10 * mm, pagesize[0] - d.rightMargin, 10 * mm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(grey)
        canvas.drawString(d.leftMargin, 6.5 * mm, f"{BRAND} - {TAGLINE}  |  Laporan Penjualan {periode}  |  Dokumen internal, khusus Owner")
        canvas.drawRightString(pagesize[0] - d.rightMargin, 6.5 * mm, f"Halaman {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()
